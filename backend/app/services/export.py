from __future__ import annotations

import csv
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from .. import config
from ..schemas import ExportFormat, Granularity, Kind
from . import series

_VALUE = {Kind.forecast: ("prediction", "Прогноз посадок"), Kind.history: ("boardings", "Посадки (факт)")}
_TITLE = {Kind.forecast: "Прогноз", Kind.history: "История"}
_LABEL = {"route": "Маршрут", "date": "Дата", "hour": "Час", "month": "Месяц"}
_MIME = {
    ExportFormat.csv: "text/csv; charset=utf-8",
    ExportFormat.xlsx: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def _clip(store, kind: Kind, start: date, end: date) -> tuple[date, date]:
    lo, hi = series.bounds(store, kind.value)
    return max(start, lo), min(end, hi)


def filename(store, kind: Kind, fmt: ExportFormat, start: date, end: date, route: int | None, gran: Granularity) -> str:
    start, end = _clip(store, kind, start, end)
    scope = f"route{route}" if route is not None else "all"
    return f"tramway_{kind.value}_{start}_{end}_{scope}_{gran.value}.{fmt.value}"


def _csv(rows: list[dict], columns: list[str], value_col: str) -> bytes:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";", lineterminator="\n")
    w.writerow(columns[:-1] + [value_col])
    w.writerows([r[c] for c in columns] for r in rows)
    return out.getvalue().encode("utf-8")


def _xlsx(rows: list[dict], columns: list[str], value_label: str, meta: list[tuple[str, str]], per_route: dict[int, int]) -> bytes:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Данные")
    ws.append([_LABEL.get(c, value_label) for c in columns])
    for r in rows:
        ws.append([r[c] for c in columns])
    ws.column_dimensions["A"].width = 12
    ws.freeze_panes = "A2"

    info = wb.create_sheet("Сводка")
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 34
    for k, v in meta:
        info.append([k, v])
    info.append([])
    info.append(["Маршрут", value_label])
    for route, total in per_route.items():
        info.append([route, total])
    info.append(["Итого", sum(per_route.values())])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _build(store, kind: Kind, fmt: ExportFormat, start: date, end: date, route: int | None, gran: Granularity) -> bytes:
    source = store.forecast if kind is Kind.forecast else store.history
    rows = series.rows(store, source, start, end, route, gran)
    period = "month" if gran is Granularity.month else "date"
    columns = ["route", period] + (["hour"] if gran is Granularity.hour else []) + ["passengers"]
    col, label = _VALUE[kind]
    if fmt is ExportFormat.csv:
        return _csv(rows, columns, col)
    per_route: dict[int, int] = {}
    for r in rows:
        per_route[r["route"]] = per_route.get(r["route"], 0) + r["passengers"]
    a, b = _clip(store, kind, start, end)
    meta = [
        ("Отчёт", f"{_TITLE[kind]} пассажиропотока трамваев"),
        ("Период", f"{a.isoformat()} — {b.isoformat()}"),
        ("Маршрут", str(route) if route is not None else "все"),
        ("Детализация", {"hour": "по часам", "day": "по дням", "month": "по месяцам"}[gran.value]),
        ("Строк в выгрузке", str(len(rows))),
        ("Сформировано", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    return _xlsx(rows, columns, label, meta, per_route)


def build(store, kind: Kind, fmt: ExportFormat, start: date, end: date, route: int | None, gran: Granularity) -> tuple[bytes, str, str]:
    series.validate(store, kind.value, start, end, route, gran, limit_hours=False)
    body = series.cached(store, ("export", kind, fmt, start, end, route, gran), lambda: _build(store, kind, fmt, start, end, route, gran))
    return body, _MIME[fmt], filename(store, kind, fmt, start, end, route, gran)
