import pandas as pd
import pytest
from fastapi.testclient import TestClient

from backend.app import config
from backend.main import app


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    mp = pytest.MonkeyPatch()
    mp.setattr(config, "OVERLAY", tmp_path_factory.mktemp("ov") / "ingested.csv")
    with TestClient(app) as c:
        yield c
    mp.undo()


@pytest.fixture(scope="module")
def sub():
    return pd.read_csv(config.SUBMISSION, sep=";", parse_dates=["date"])


def test_health(client):
    body = client.get("/api/health").json()
    assert body["status"] == "ok" and body["forecast_rows"] == 14640 and body["routes"] == 10


def test_v1_and_legacy_alias_equal(client):
    a = client.get("/api/forecast?start=2025-11-01&end=2025-11-03").json()
    b = client.get("/api/v1/forecast?start=2025-11-01&end=2025-11-03").json()
    assert a == b and len(a) == 30


def test_forecast_matches_submission(client, sub):
    rows = client.get("/api/forecast?granularity=hour&route=7&start=2025-11-10&end=2025-11-12").json()
    exp = sub[(sub.route == 7) & (sub.date >= "2025-11-10") & (sub.date <= "2025-11-12")]
    assert len(rows) == 72 and sum(r["passengers"] for r in rows) == int(exp.prediction.round().sum())


def test_day_and_month_totals_agree(client, sub):
    day = client.get("/api/forecast?route=1").json()
    month = client.get("/api/forecast?route=1&granularity=month").json()
    assert sum(r["passengers"] for r in day) == sum(r["passengers"] for r in month)
    assert [r["month"] for r in month] == ["2025-11", "2025-12"]
    assert sum(r["passengers"] for r in day) == int(sub[sub.route == 1].prediction.round().sum())


def test_history_route5_is_zero(client):
    assert sum(r["passengers"] for r in client.get("/api/history?route=5").json()) == 0


def test_history_total_matches_labels(client):
    labels = pd.concat([pd.read_csv(p, sep=";") for p in config.LABELS])
    got = sum(r["passengers"] for r in client.get("/api/history").json())
    assert got == int(labels.boardings.sum())


def test_weekday_average(client):
    rows = client.get("/api/history/weekday-average?route=1").json()
    assert len(rows) == 7 and all(r["route"] == 1 for r in rows)


@pytest.mark.parametrize(
    "url,status,code",
    [
        ("/api/forecast?start=2025-12-01&end=2025-11-01", 400, "bad_range"),
        ("/api/forecast?route=99", 404, "route_not_found"),
        ("/api/forecast?start=2024-01-01&end=2024-01-05", 422, "out_of_range"),
        ("/api/forecast?granularity=hour", 422, "range_too_wide"),
        ("/api/forecast?granularity=week", 422, "validation_error"),
        ("/api/forecast?start=вчера", 422, "validation_error"),
        ("/api/history?route=abc", 422, "validation_error"),
        ("/api/history/weekday-average?route=99", 404, "route_not_found"),
    ],
)
def test_errors_are_clear(client, url, status, code):
    r = client.get(url)
    assert r.status_code == status and r.json()["code"] == code and r.json()["detail"]


def test_map_and_export(client):
    assert client.get("/api/map").json()["type"] == "FeatureCollection"
    assert client.get("/api/export").text.startswith("route;date;hour;prediction")


def test_export_csv_default_equals_submission(client, sub):
    r = client.get("/api/export")
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert "attachment" in r.headers["content-disposition"]
    lines = r.text.strip().split("\n")
    assert lines[0] == "route;date;hour;prediction" and len(lines) == 14641
    assert sum(int(x.split(";")[3]) for x in lines[1:]) == int(sub.prediction.round().sum())


def test_export_respects_filters(client):
    r = client.get("/api/export?route=7&start=2025-11-01&end=2025-11-07&granularity=day")
    lines = r.text.strip().split("\n")
    assert lines[0] == "route;date;prediction" and len(lines) == 8 and all(l.startswith(("route", "7;")) for l in lines)
    assert "route7" in r.headers["content-disposition"]


def test_export_history_month_and_wide_hours_allowed(client):
    assert client.get("/api/export?kind=history&granularity=month").text.split("\n")[0] == "route;month;boardings"
    assert len(client.get("/api/export?kind=history&granularity=hour").text.strip().split("\n")) == 10 * 304 * 24 + 1


def test_export_xlsx(client):
    import io

    from openpyxl import load_workbook

    r = client.get("/api/export?format=xlsx&route=1&granularity=day")
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Данные", "Сводка"]
    assert wb["Данные"].max_row == 62 and wb["Данные"]["A1"].value == "Маршрут"
    assert wb["Сводка"]["B2"].value.startswith("2025-11-01")


@pytest.mark.parametrize("url,status", [("/api/export?route=99", 404), ("/api/export?format=pdf", 422), ("/api/export?kind=x", 422), ("/api/export?start=2030-01-01&end=2030-01-02", 422)])
def test_export_errors(client, url, status):
    assert client.get(url).status_code == status


def test_adjusted_neutral_equals_base(client):
    r = client.post("/api/v1/forecast/adjusted", json={"route": 7, "granularity": "day"}).json()
    assert r["summary"]["delta"] == 0 and all(p["passengers"] == p["base"] for p in r["data"])


def test_adjusted_global_and_rules(client):
    r = client.post("/api/v1/forecast/adjusted", json={"route": 7, "factors": {"weather": 0.9, "event": 1.1}}).json()
    assert r["summary"]["global_multiplier"] == 0.99 and abs(r["summary"]["delta_pct"] + 1) < 0.1
    body = {"route": 7, "granularity": "hour", "start": "2025-11-05", "end": "2025-11-05",
            "rules": [{"start": "2025-11-05", "end": "2025-11-05", "factor": 2.0, "hour_from": 8, "hour_to": 9}]}
    rows = client.post("/api/v1/forecast/adjusted", json=body).json()["data"]
    for p in rows:
        want = p["base"] * 2 if p["hour"] in (8, 9) else p["base"]
        assert p["passengers"] == want


def test_adjusted_rule_scoped_to_routes(client):
    body = {"granularity": "day", "start": "2025-11-05", "end": "2025-11-05",
            "rules": [{"start": "2025-11-05", "end": "2025-11-05", "factor": 0.5, "routes": [1]}]}
    rows = {p["route"]: p for p in client.post("/api/v1/forecast/adjusted", json=body).json()["data"]}
    assert abs(rows[1]["passengers"] - rows[1]["base"] * 0.5) <= 12 and rows[7]["passengers"] == rows[7]["base"]


@pytest.mark.parametrize(
    "body,status,code",
    [
        ({"factors": {"weather": 3}}, 422, "validation_error"),
        ({"rules": [{"start": "2025-11-09", "end": "2025-11-01", "factor": 1}]}, 422, "validation_error"),
        ({"rules": [{"start": "2025-11-01", "end": "2025-11-02", "factor": 1, "routes": [99]}]}, 404, "route_not_found"),
        ({"rules": [{"start": "2024-01-01", "end": "2024-01-02", "factor": 1}]}, 422, "rule_out_of_range"),
        ({"route": 99}, 404, "route_not_found"),
        ({"granularity": "hour"}, 422, "range_too_wide"),
    ],
)
def test_adjusted_errors(client, body, status, code):
    r = client.post("/api/v1/forecast/adjusted", json=body)
    assert r.status_code == status and r.json()["code"] == code and r.json()["detail"]


def test_year_outlook(client):
    r = client.get("/api/v1/forecast/year?route=1&growth=1.1").json()
    by = {(p["month"]): p for p in r["data"]}
    assert len(r["data"]) == 24 and by["2025-01"]["source"] == "fact" and by["2025-11"]["source"] == "forecast"
    assert by["2026-11"]["source"] == "scenario"
    assert abs(by["2026-11"]["passengers"] - by["2025-11"]["passengers"] * 1.1) <= 1  # в ноябре по 30 дней
    assert client.get("/api/forecast/year?growth=9").status_code == 422


def test_factors_presets_are_measured_and_consistent(client):
    f = client.get("/api/factors").json()
    assert f["limits"]["factor"] == [0.5, 1.5] and f["sources"]["calendar"].startswith("https://")
    rain = next(p for p in f["weather"] if p.get("measured"))
    assert rain["value"] < 1 and rain["ci95"][0] < rain["value"] < rain["ci95"][1]  # дождь снижает посадки, а не повышает
    assert f["calendar"]["days"] == ["2025-11-03", "2025-11-04", "2025-12-31"] and 0.4 < f["calendar"]["holiday_factor"] < 0.5


def test_calendar_endpoint(client):
    body = client.get("/api/v1/calendar?start=2025-11-01&end=2025-11-05").json()
    types = {d["date"]: d["type"] for d in body["days"]}
    assert types["2025-11-03"] == "holiday" and types["2025-11-05"] == "work" and types["2025-11-02"] == "weekend"
    assert next(d for d in body["days"] if d["date"] == "2025-11-04")["factor"] == body["holiday_factor"]
    assert client.get("/api/v1/calendar?start=2030-01-01&end=2030-01-02").status_code == 422
    assert client.get("/api/v1/calendar?start=2025-12-01&end=2025-11-01").status_code == 400


def test_adjusted_calendar_does_not_double_count_ml_holidays(client):
    body = {"route": 7, "granularity": "day", "start": "2025-11-02", "end": "2025-11-05", "calendar": True}
    r = client.post("/api/v1/forecast/adjusted", json=body).json()
    by = {p["date"]: p for p in r["data"]}
    assert all(row["passengers"] == row["base"] for row in by.values())
    assert r["summary"]["calendar_in_model"] is True
    assert r["summary"]["calendar_days"] == ["2025-11-03", "2025-11-04", "2025-12-31"]
    assert client.post("/api/v1/forecast/adjusted", json={"route": 7}).json()["summary"]["calendar_days"] == []


def test_calendar_unavailable_is_reported(client, tmp_path, monkeypatch):
    from backend.app.store import DataStore

    monkeypatch.setattr(config, "EXTERNAL_DIR", tmp_path)
    client.app.state.store = DataStore()
    try:
        r = client.post("/api/v1/forecast/adjusted", json={"calendar": True})
        assert r.status_code == 422 and r.json()["code"] == "calendar_unavailable"
    finally:
        monkeypatch.undo()
        client.app.state.store = DataStore()
