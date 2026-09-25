"""Справочник xlsx → artifacts/stops.json: остановки маршрутов в порядке следования с координатами и районом.

    python -m backend.build_stops
"""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
workbooks = list((ROOT / "dataset" / "spravochniki").glob("*справочники*.xlsx"))
if not workbooks:
    raise FileNotFoundError("Не найден справочник трамвайных маршрутов")
book = openpyxl.load_workbook(workbooks[0], read_only=True, data_only=True)

# stop_id → (название, район)
districts = {str(r[0]): (r[1], r[7]) for r in book["Остановки GTFS_STOPS"].iter_rows(min_row=3, values_only=True) if r[0]}

trips: dict[tuple[int, int, str], list[tuple]] = defaultdict(list)
for row in book["Порядок_с_координатами"].iter_rows(min_row=3, values_only=True):
    route, trip, direction, seq, stop_id, name, lat, lon = row[1], row[4], row[6], row[9], row[10], row[14], row[15], row[16]
    if not (route and stop_id and lat and lon):
        continue
    trips[(int(route), int(direction), str(trip))].append((int(seq), str(stop_id), name, float(lat), float(lon)))

# На направление берём самый длинный вариант рейса
best: dict[tuple[int, int], list[tuple]] = {}
for (route, direction, _), stops in trips.items():
    if len(stops) > len(best.get((route, direction), [])):
        best[(route, direction)] = sorted(stops)

served: dict[str, set[int]] = defaultdict(set)
for (route, _), stops in best.items():
    for _, stop_id, *_ in stops:
        served[stop_id].add(route)

records = []
for (route, direction), stops in sorted(best.items()):
    for seq, stop_id, name, lat, lon in stops:
        ref_name, district = districts.get(stop_id, (name, None))
        records.append({
            "route": route, "direction": direction, "sequence": seq, "stop_id": stop_id,
            "name": name or ref_name or "Остановка", "lat": lat, "lon": lon, "district": district,
            "routes_serving": sorted(served[stop_id]),
        })

out = ROOT / "artifacts" / "stops.json"
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
print(f"Сохранено {len(records)} остановок ({len(best)} направлений, маршруты {sorted({r for r, _ in best})}) в {out}")
